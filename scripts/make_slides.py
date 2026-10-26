"""Module 1 five-slide deck -> presentations/module1_presentation.pptx (numbers from executed results)."""
from __future__ import annotations

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from deliverables_common import FIG, REPO_URL, ROOT, hypotheses, load

NAVY = RGBColor(0x0B, 0x2A, 0x4A)
BLUE = RGBColor(0x2F, 0x6F, 0xAE)
ICE = RGBColor(0xEA, 0xF2, 0xFA)
EDGE = RGBColor(0xC9, 0xDB, 0xEE)
GREY = RGBColor(0x5F, 0x6B, 0x76)
ACCENT = RGBColor(0xE0, 0x7B, 0x39)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SOFT = RGBColor(0xCF, 0xE2, 0xF3)
GREEN = RGBColor(0x1F, 0x6F, 0x2A)
FONT = "Calibri"
KICK = "MODULE 1 · IMAGE-BASED VISION TRANSFORMER"


def text(slide, x, y, w, h, runs, size=16, color=NAVY, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, name=None):
    """runs: str or list of paragraphs; each paragraph is str or list of (text, {bold,color,size}) tuples."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    paras = runs if isinstance(runs, list) else [runs]
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        for t, o in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", bold)
            r.font.color.rgb = o.get("color", color)
    return tb


def bullets(slide, x, y, w, h, items, size=15, color=NAVY, space=6):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        parts = it if isinstance(it, list) else [(it, {})]
        r0 = p.add_run()
        r0.text = "▪  "
        r0.font.size, r0.font.color.rgb, r0.font.name = Pt(size), BLUE, FONT
        for t, o in parts:
            r = p.add_run()
            r.text = t
            r.font.name = FONT
            r.font.size = Pt(o.get("size", size))
            r.font.bold = o.get("bold", False)
            r.font.color.rgb = o.get("color", color)
    return tb


def card(slide, x, y, w, h, fill=ICE, line=EDGE, name=None):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    s.adjustments[0] = 0.06
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if name:
        s.name = name
    return s


def header(slide, prs, title, num, dark=False):
    c = SOFT if dark else BLUE
    text(slide, Inches(0.6), Inches(0.35), Inches(10), Inches(0.35), KICK, size=12, color=c, bold=True)
    text(slide, Inches(0.6), Inches(0.68), Inches(12.1), Inches(0.8), title, size=32, color=WHITE if dark else NAVY, bold=True)
    text(slide, Inches(0.6), prs.slide_height - Inches(0.45), Inches(10.5), Inches(0.3),
         "Proof of concept on public simulation data (openEMS, RadPat-50K) · Ansys HFSS not run · results not HFSS-validated",
         size=10, color=SOFT if dark else GREY)
    text(slide, prs.slide_width - Inches(1.4), prs.slide_height - Inches(0.45), Inches(0.8), Inches(0.3), f"{num} / 5",
         size=10, color=SOFT if dark else GREY, align=PP_ALIGN.RIGHT)


def background(slide, prs, color):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = color


def stat(slide, x, y, w, big, label, color=NAVY, sub=GREY, size=40):
    text(slide, x, y, w, Inches(0.75), big, size=size, color=color, bold=True)
    text(slide, x, y + Inches(0.72), w, Inches(0.6), label, size=13, color=sub)


def table(slide, x, y, w, rows, col_w, size=14, bold_rows=(), row_h=0.38):
    shp = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(row_h * len(rows)))
    t = shp.table
    for j, cw in enumerate(col_w):
        t.columns[j].width = cw
    for i, r in enumerate(rows):
        t.rows[i].height = Inches(row_h)
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = ""
            p = c.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = str(v)
            run.font.size, run.font.name = Pt(size), FONT
            run.font.bold = i == 0 or i in bold_rows
            run.font.color.rgb = WHITE if i == 0 else NAVY
            p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
            c.fill.solid()
            c.fill.fore_color.rgb = NAVY if i == 0 else (ICE if i % 2 else WHITE)
            c.margin_top = c.margin_bottom = Emu(25000)
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
    return t


def main():
    R = load()
    row, fmt = R["row"], R["fmt"]
    vit, cnn = row("patch", "perf_class", "finetune"), row("patch", "perf_class", "cnn")
    lp, maj = row("patch", "perf_class", "linear_probe"), row("patch", "perf_class", "majority")
    vb, lb = row("patch", "band_class", "finetune"), row("patch", "band_class", "linear_probe")
    vr = row("radpat", "hpbw_class", "finetune")
    vood = row("patch", "perf_class", "finetune", "ood")
    noa = row("patch", "perf_class", "finetune_noaug")
    e, st, prep = R["expl"], R["stats"], R["prep"]
    sc = prep["patch_split_counts"]
    H = hypotheses(R)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W = prs.slide_width

    # ---------------------------------------------------------------- 1 problem and objective (dark)
    s = prs.slides.add_slide(blank)
    background(s, prs, NAVY)
    header(s, prs, "Can a ViT read antenna performance from an image?", 1, dark=True)
    text(s, Inches(0.6), Inches(1.5), Inches(7.4), Inches(0.9),
         "Vision Transformer-based visual analysis of hybrid-slotted DGS compact patch antennas for 5G and emerging 6G networks",
         size=16, color=SOFT)
    card(s, Inches(0.6), Inches(2.45), Inches(7.4), Inches(4.3), fill=RGBColor(0x14, 0x3A, 0x62), line=None)
    text(s, Inches(0.85), Inches(2.6), Inches(7.0), Inches(0.4), "Problem, gap and objective", size=18, color=WHITE, bold=True)
    bullets(s, Inches(0.85), Inches(3.1), Inches(7.0), Inches(3.6), [
        "Slots and DGS reshape surface currents: small geometry changes shift resonance and matching",
        "Full-wave sweeps (HFSS) are slow; most ML antenna studies use tables only",
        "Gap: little evidence that standardised images + a ViT can screen performance and explain why",
        "Objective: build and test a reproducible, explainable image classifier for antenna performance",
        "Scope: no public HFSS DGS image set exists and HFSS was unavailable → proof of concept on open simulations",
    ], size=15, color=WHITE)
    card(s, Inches(8.3), Inches(1.5), Inches(4.45), Inches(5.25), fill=WHITE, line=None)
    text(s, Inches(8.55), Inches(1.65), Inches(4.0), Inches(0.4), "Pre-registered hypotheses", size=18, color=NAVY, bold=True)
    hyp = ["H1  ViT beats the majority baseline", "H2  ViT ≥ compact CNN (3 seeds)", "H3  Fine-tuning beats a frozen probe",
           "H4  Attention concentrates on patch and feed", "H5  Scores drop on unseen shapes (OOD)"]
    bullets(s, Inches(8.55), Inches(2.15), Inches(4.0), Inches(2.6), hyp, size=15)
    s.shapes.add_picture(str(ROOT / "data" / "processed" / "patch_images" / "P1307.png"), Inches(8.55), Inches(4.75), height=Inches(1.85))
    text(s, Inches(10.5), Inches(4.85), Inches(2.1), Inches(1.7),
         [[("Real model input", {"bold": True, "size": 13})], "Design P1307, class High. Top: true scale. Bottom: zoomed, so the feed dot is resolvable."],
         size=12, color=GREY)
    s.notes_slide.notes_text_frame.text = (
        "Module 1 is the image branch of the two-module framework. The question: can a Vision Transformer look at a standardised "
        "picture of an antenna and predict how well it is matched, and can it show where it looked? HFSS was not available and no "
        "licensed HFSS image dataset of hybrid-slotted DGS antennas exists, so this is a proof of concept on public simulation data. "
        "The five hypotheses were committed to Git before any fine-tuned result existed. Reference: Dosovitskiy et al., ICLR 2021.")

    # ---------------------------------------------------------------- 2 data and preprocessing
    s = prs.slides.add_slide(blank)
    background(s, prs, WHITE)
    header(s, prs, "Two verified public datasets, one standard image per design", 2)
    for i, (big, lab) in enumerate([("2,051", "openEMS patch designs\nHugging Face · MIT"), ("3,000", "RadPat-50K polar plots\nZenodo · CC BY 4.0"),
                                    ("70 / 15 / 15", f"{sc['train']:,} · {sc['val']} · {sc['test']} designs\ndesign-level split, seed 42"),
                                    (str(sc["ood"]), "OOD hold-out\nwidest patches, W/L > 1.95")]):
        x = Inches(0.6 + i * 3.1)
        card(s, x, Inches(1.6), Inches(2.85), Inches(1.55))
        stat(s, x + Inches(0.2), Inches(1.68), Inches(2.6), big, lab, size=36)
    text(s, Inches(0.6), Inches(3.4), Inches(5.6), Inches(0.4), "Labels from simulated S11 (frozen before training)", size=17, bold=True)
    bullets(s, Inches(0.6), Inches(3.85), Inches(5.7), Inches(3.0), [
        [("Matching: ", {"bold": True}), ("Poor > −10 dB · Acceptable −20 to −10 dB · High ≤ −20 dB*", {})],
        [("Band: ", {"bold": True}), ("first matched resonance < 5 | 5–7.125 | ≥ 7.125 GHz (3GPP)", {})],
        [("Beamwidth (RadPat): ", {"bold": True}), ("training-split HPBW terciles", {})],
        [("Leakage control: ", {"bold": True}), ("no text or legends, MD5 de-duplication, one image per design", {})],
        [("Augmentation (train only): ", {"bold": True}), ("mirror, ±3° rotation, 2% shift", {})],
    ], size=14)
    text(s, Inches(0.6), Inches(6.5), Inches(5.6), Inches(0.4), "*−20 dB is a project assumption that needs engineering confirmation",
         size=11, color=GREY)
    s.shapes.add_picture(str(FIG / "dataset_overview.png"), Inches(6.45), Inches(3.45), width=Inches(6.3))
    s.shapes.add_picture(str(FIG / "radpat_samples.png"), Inches(6.9), Inches(5.55), width=Inches(5.4))
    s.notes_slide.notes_text_frame.text = (
        "Datasets were searched in the order the brief sets: Hugging Face, Kaggle, Mendeley, IEEE DataPort, GitHub, plus Zenodo. "
        "The openEMS patch sweep gives geometry and full S11 curves; each design is drawn as a two-panel image with fixed colours "
        "and no text. RadPat-50K adds a radiation-pattern modality. The tabular Mendeley data were assigned to Module 2. "
        "References: LaBash et al. 2025 (arXiv:2505.18188); Bhuma et al. 2025 (doi:10.5281/zenodo.17067881).")

    # ---------------------------------------------------------------- 3 architecture
    s = prs.slides.add_slide(blank)
    background(s, prs, WHITE)
    header(s, prs, "ViT-S/16 pipeline with three baselines and explainability", 3)
    s.shapes.add_picture(str(FIG / "module1_architecture.png"), Inches(0.6), Inches(1.5), width=Inches(9.2))
    card(s, Inches(10.05), Inches(1.5), Inches(2.7), Inches(4.4))
    text(s, Inches(10.2), Inches(1.6), Inches(2.45), Inches(0.4), "Training recipe", size=16, bold=True)
    bullets(s, Inches(10.2), Inches(2.05), Inches(2.5), Inches(3.8), [
        "ImageNet-21k pretrained, 196 tokens + CLS",
        "Stage A: head only, 2 epochs",
        "Stage B: last 4 of 12 blocks, lr 3e-5",
        "AdamW, batch 16, early stop on val macro-F1",
        "Temperature-scaled probabilities",
        "CPU only (i7-10610U)",
    ], size=14, space=4)
    text(s, Inches(0.6), Inches(6.15), Inches(12.1), Inches(0.7),
         [[("Outputs: ", {"bold": True}), ("matching / band / beamwidth class with calibrated confidence, attention-rollout maps, and a 384-d "
                                           "embedding per design that Module 2 uses for multimodal fusion. Dashed boxes need HFSS.", {})]],
         size=14)
    s.notes_slide.notes_text_frame.text = (
        "Each 224 by 224 image is cut into 196 patches of 16 by 16 pixels. Self-attention lets the feed dot relate to the patch "
        "edges anywhere in the image. Three comparators run on the same split: majority class, a compact CNN with 0.25 million "
        "parameters, and a frozen-ViT linear probe. Explainability uses attention rollout (Abnar and Zuidema, 2020), scored "
        "against analytic patch and feed masks. Deviations for the CPU budget: 4 of 12 blocks unfrozen, patience 3.")

    # ---------------------------------------------------------------- 4 results
    s = prs.slides.add_slide(blank)
    background(s, prs, WHITE)
    header(s, prs, "All models beat chance; augmentation held the ViT back", 4)
    table(s, Inches(0.6), Inches(1.55), Inches(6.2), [
        ["Matching class, locked test (n = 262)", "Macro-F1", "Bal. acc."],
        ["Majority baseline", f"{maj.macro_f1:.3f}", f"{maj.balanced_accuracy:.3f}"],
        ["Compact CNN (3 seeds)", fmt(cnn), f"{cnn.balanced_accuracy:.3f}"],
        ["ViT-S/16 frozen linear probe", f"{lp.macro_f1:.3f}", f"{lp.balanced_accuracy:.3f}"],
        ["ViT-S/16 fine-tuned (3 seeds)", fmt(vit), f"{vit.balanced_accuracy:.3f}"],
        ["ViT fine-tuned, no augmentation*", f"{noa.macro_f1:.3f}", f"{noa.balanced_accuracy:.3f}"],
    ], [Inches(3.5), Inches(1.6), Inches(1.1)], size=14, bold_rows=(4,))
    card(s, Inches(0.6), Inches(4.05), Inches(6.2), Inches(2.45))
    p_vc = st["mcnemar_finetune_vs_cnn_perf_class_s42"]["exact_p_value"]
    bullets(s, Inches(0.8), Inches(4.15), Inches(5.9), Inches(2.6), [
        [("H1 supported; H2, H3 not: ", {"bold": True}), (f"CNN beat the ViT (McNemar p = {p_vc:.4f})", {})],
        [("Why: ", {"bold": True}), (f"±3° / 2% shifts blur the 1–2 px feed offset; without them {vit.macro_f1:.2f} → {noa.macro_f1:.2f}", {})],
        [("H4: ", {"bold": True}), (f"feed gets ×{e['median_feed_enrichment']:.2f} its area share, patch ×{e['median_patch_enrichment']:.2f}", {})],
        [("H5 not supported: ", {"bold": True}), (f"OOD {vood.macro_f1:.2f}, confounded by class balance", {})],
        [("Other tasks (ViT): ", {"bold": True}), (f"band {vb.macro_f1:.2f} (probe {lb.macro_f1:.2f}), beamwidth {vr.macro_f1:.2f}", {})],
    ], size=14, space=4)
    s.shapes.add_picture(str(FIG / "model_comparison.png"), Inches(7.05), Inches(1.5), width=Inches(5.7))
    s.shapes.add_picture(str(FIG / "attention_rollout_examples.png"), Inches(7.05), Inches(4.0), width=Inches(5.7))
    text(s, Inches(7.05), Inches(6.55), Inches(5.7), Inches(0.3),
         "*exploratory, one seed · attention maps are model attribution, not EM causation", size=11, color=GREY)
    s.notes_slide.notes_text_frame.text = (
        "All values come from saved metric files of executed runs. Hypothesis outcomes: "
        + "; ".join(f"{h[0]} {h[2]}" for h in H)
        + ". The training curves show underfitting. The ablation planned in the Module 1 report found the cause: the "
          "augmentation blurred the feed offset that sets matching. Robustness: blur and grey recolouring hurt most.")

    # ---------------------------------------------------------------- 5 conclusion (dark)
    s = prs.slides.add_slide(blank)
    background(s, prs, NAVY)
    header(s, prs, "A working, reproducible image module, ready for HFSS data", 5, dark=True)
    cols = [("Contribution", [
        "End-to-end ViT pipeline executed on CPU",
        "Geometry images carry a learnable signal for matching and band class",
        "Attribution measured against known patch and feed regions",
        "Lesson: augmentations must preserve fine physical detail"]),
        ("Limitations", [
            "Plain rectangular patches; no slots or DGS",
            "Simulation only (openEMS, analytical patterns)",
            "−20 dB boundary assumed; CPU-limited tuning",
            "Attention ≠ causation; OOD test confounded"]),
        ("Next: HFSS validation", [
            "≥ 1,000 hybrid-slotted DGS designs via PyAEDT",
            "Geometry, current and E-field images at fixed scale",
            "Retrain unchanged; hold out a slot topology",
            "Verify screened designs in full-wave"])]
    for i, (t, items) in enumerate(cols):
        x = Inches(0.6 + i * 4.1)
        card(s, x, Inches(1.6), Inches(3.85), Inches(3.45), fill=RGBColor(0x14, 0x3A, 0x62), line=None)
        text(s, x + Inches(0.2), Inches(1.72), Inches(3.5), Inches(0.4), t, size=18, color=WHITE, bold=True)
        bullets(s, x + Inches(0.2), Inches(2.2), Inches(3.5), Inches(3.1), items, size=15, color=WHITE)
    card(s, Inches(0.6), Inches(5.4), Inches(12.15), Inches(0.95), fill=WHITE, line=None)
    text(s, Inches(0.85), Inches(5.5), Inches(11.7), Inches(1.0), [
        [("Deliverables: ", {"bold": True}), ("code, configs, split manifests, 18 passing tests, data and model cards, Word/PDF/LaTeX report, this deck, poster", {})],
        [("GitHub: ", {"bold": True}), (REPO_URL, {"color": BLUE})]], size=14)
    s.notes_slide.notes_text_frame.text = (
        "Module 1 shows the pipeline works end to end, from rendering through training, testing and explaining. It does not "
        "show anything yet about hybrid-slotted DGS antennas; that needs the HFSS dataset described in the Module 1 report. "
        "The image embeddings feed Module 2's fusion experiment. References: Dosovitskiy et al. (2021); Abnar and Zuidema (2020); "
        "Guo et al. (2017); Darcet et al. (2024); 3GPP TS 38.101-1.")

    out = ROOT / "presentations"
    out.mkdir(exist_ok=True)
    prs.save(out / "module1_presentation.pptx")
    print("slides written")


if __name__ == "__main__":
    main()
